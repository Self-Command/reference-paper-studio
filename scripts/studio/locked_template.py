"""Versioned mother documents and a content-only Word writer. Actions only."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import json
import os
import re
import zipfile
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips, RGBColor
from docx.enum.text import WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from .runtime import require_actions, StudioError, digest, write_json
from .docx_layout import W, M, section_break, is_equation_table
from .vendor.ptlib.apply import _apply_sectpr, _three_line
from .vendor.ptlib.ooxml import set_child, set_fonts

ROOT = Path(__file__).resolve().parents[2]
POLICY = ROOT / "assets/templates/reference-journal-v1.json"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
ROLES = ("Body", "Title", "Author", "Abstract", "Keywords", "Caption", "TableText", "Equation", "Marker", "Heading1", "Heading2", "Heading3", "Heading4", "List")


def policy():
    require_actions()
    return json.loads(POLICY.read_text(encoding="utf-8"))


def resolve_template(request):
    require_actions()
    if request.get("schema_version") == 2:
        from .contracts import validate
        from .registry import resolve
        validate(request,"request")
        return resolve(request)
    forbidden = {"profile", "fonts", "font", "font_size", "body_size_pt", "page", "margins", "columns", "sections", "template_path", "styles", "layout"}
    extra = forbidden.intersection(request)
    if extra: raise StudioError("TEMPLATE_OVERRIDE_FORBIDDEN: " + ", ".join(sorted(extra)))
    allowed = {"id", "stage", "files", "sha256", "manuscript", "metadata", "bibliography", "reference_evidence", "model_evidence", "document_kind", "expected_equations", "template_id", "template_version", "template_variant"}
    if set(request)-allowed: raise StudioError("UNKNOWN_REQUEST_FIELDS: "+str(sorted(set(request)-allowed)))
    p = policy()
    if request.get("template_id") != p["id"] or request.get("template_version") != p["version"]:
        raise StudioError("UNKNOWN_TEMPLATE_VERSION: require reference-journal-v1 version 1")
    variant = request.get("template_variant", "front")
    if variant not in ("front", "end"): raise StudioError("UNKNOWN_TEMPLATE_VARIANT")
    directory = Path(os.environ.get("RPS_TEMPLATE_ROOT", ROOT/"assets/templates/v1"))
    manifest = json.loads((directory/"manifest.json").read_text(encoding="utf-8"))
    if manifest["policy_sha256"] != digest(POLICY): raise StudioError("TEMPLATE_POLICY_FINGERPRINT_MISMATCH")
    path = directory/manifest["variants"][variant]["file"]
    if digest(path) != manifest["variants"][variant]["sha256"]: raise StudioError("TEMPLATE_FINGERPRINT_MISMATCH")
    p["english_abstract_position"] = variant
    p["equation_numbering"] = "sequential" if request.get("document_kind") == "paper" else "preserve"
    return path, p, {"id": p["id"], "version": p["version"], "variant": variant, "sha256": digest(path), "policy_sha256": digest(POLICY)}


def column_width(p):
    margins = p["page"]["margin_cm"]
    columns=p.get("body_columns",2)
    page_width=p["page"].get("width_cm",21)
    return round((page_width-margins["left"]-margins["right"]-(columns-1)*p["page"]["column_spacing_cm"])/columns*1440/2.54)


def pstyle(element, role):
    old = element.find(W+"pPr")
    numbering = deepcopy(old.find(W+"numPr")) if old is not None and old.find(W+"numPr") is not None else None
    if old is not None: element.remove(old)
    props = OxmlElement("w:pPr"); element.insert(0, props)
    set_child(props, "pStyle", {"w:val": "RPS"+role})
    if numbering is not None: props.append(numbering)


def paragraph(text, role):
    element = OxmlElement("w:p"); pstyle(element, role)
    run = OxmlElement("w:r"); node = OxmlElement("w:t"); node.text = text
    run.append(node); element.append(run)
    return element


def create_templates(directory, policy_override=None):
    """Create release assets once; never build a mother from a manuscript."""
    require_actions()
    directory = Path(directory); directory.mkdir(parents=True, exist_ok=True)
    p = policy_override or policy(); variants = {}
    for variant in ("front", "end"):
        doc = Document(); body = doc._element.body
        for item in list(body):
            if item.tag != W+"sectPr": body.remove(item)
        for role in ROLES:
            style = doc.styles.add_style("RPS"+role, 1); style.base_style = doc.styles["Normal"]
            style.font.name = p["fonts"]["latin"]
            style.font.size = Pt(16 if role == "Title" else 12 if role == "Heading1" else 1 if role == "Marker" else p["body_size_pt"])
            style.font.bold = role == "Title" or role.startswith("Heading")
            style.font.color.rgb = RGBColor(0, 0, 0)
            set_fonts(style.element.get_or_add_rPr(), latin=p["fonts"]["latin"], cjk=p["fonts"]["cjk"])
            fmt = style.paragraph_format
            fmt.space_before = Pt(6 if role.startswith("Heading") else 0)
            fmt.space_after = Pt(0 if role == "Marker" else 2 if role == "TableText" else 3)
            fmt.line_spacing = Pt(1 if role == "Marker" else 24 if role == "Title" else p.get("table_line_minimum_pt",12) if role == "TableText" else p.get("body_line_minimum_pt",14))
            fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY if role == "Marker" else WD_LINE_SPACING.AT_LEAST
            fmt.keep_with_next = role.startswith("Heading"); fmt.keep_together = False
            fmt.widow_control = True; fmt.page_break_before = False
            fmt.first_line_indent = Pt(20 if role == "Body" else 0)
            fmt.left_indent = Pt(0); fmt.right_indent = Pt(0)
            fmt.alignment = 1 if role in ("Title", "Author", "Caption") else 3 if role in ("Body", "Abstract") else 0
        normal = doc.styles["Normal"]; normal.font.name = p["fonts"]["latin"]; normal.font.size = Pt(p["body_size_pt"])
        set_fonts(normal.element.get_or_add_rPr(), latin=p["fonts"]["latin"], cjk=p["fonts"]["cjk"])
        settings = doc.settings.element
        for inherited in list(settings.findall(M+"mathPr")): settings.remove(inherited)
        math = OxmlElement("m:mathPr"); font = OxmlElement("m:mathFont"); font.set(qn("m:val"), p["fonts"]["math"]); math.append(font); settings.append(math)
        wrap_indent=OxmlElement("m:wrapIndent"); wrap_indent.set(M+"val","0"); math.append(wrap_indent)
        def metadata(suffix):
            if suffix not in p.get("metadata_languages",["zh","en"]): return
            for key, role in (("title", "Title"), ("authors", "Author"), ("affiliation", "Author"), ("abstract", "Abstract"), ("keywords", "Keywords")):
                if key not in p.get("front_roles",["title","authors","affiliation","abstract","keywords"]): continue
                if p.get("anonymous") and key in ("authors","affiliation"): continue
                body.insert(len(body)-1, paragraph("{{"+key+"_"+suffix+"}}", role))
        metadata("zh")
        if variant == "front": metadata("en")
        body.insert(len(body)-1, section_break(1, p))
        body.insert(len(body)-1, paragraph("{{RPS_BODY}}", "Body"))
        if p.get("front_separate_page"):
            boundary=body[-2].getprevious()
            if boundary is not None and boundary.find(W+"pPr/"+W+"sectPr") is not None:
                set_child(boundary.find(W+"pPr/"+W+"sectPr"),"type",{"w:val":"nextPage"})
        if variant == "end": body.insert(len(body)-1, section_break(p.get("body_columns",2), p)); metadata("en")
        final = body.find(W+"sectPr")
        _apply_sectpr(final, {**p["page"], "columns": p.get("body_columns",2) if variant == "front" else 1, "doc_grid": False})
        set_child(final, "type", {"w:val": "continuous"})
        path = directory/("master-"+variant+".docx"); doc.save(path)
        variants[variant] = {"file": path.name, "sha256": digest(path)}
    write_json(directory/"manifest.json", {"id": p["id"], "version": p["version"], "policy_sha256": digest(POLICY), "variants": variants})


def clean_content(element, p):
    """Input cannot carry paragraph geometry or Word style definitions."""
    mapping = {"Title": "Heading1", "BodyText": "Body", "FirstParagraph": "Body", "Compact": "Body", "Caption": "Caption", "ImageCaption": "Caption", "TableCaption": "Caption"}
    for par in element.iter(W+"p"):
        old = par.find(W+"pPr/"+W+"pStyle")
        identifier = old.get(W+"val", "BodyText") if old is not None else "BodyText"
        role = identifier if identifier in {"Heading1", "Heading2", "Heading3", "Heading4"} else mapping.get(identifier, "Body")
        pstyle(par, role)
    for run in element.iter(W+"r"):
        props = run.find(W+"rPr")
        if props is not None:
            for child in list(props):
                if child.tag not in {W+"b", W+"i", W+"vertAlign"}: props.remove(child)
    for br in list(element.iter(W+"br")):
        if br.get(W+"type") in ("page", "column"): br.getparent().remove(br)
    for run in element.iter(M+"r"):
        props = run.find(W+"rPr")
        if props is not None: run.remove(props)
        props = OxmlElement("w:rPr"); run.insert(1 if run.find(M+"rPr") is not None else 0, props)
        set_fonts(props, latin=p["fonts"]["math"], cjk=p["fonts"]["cjk"])
        set_child(props, "sz", {"w:val": str(round(p["body_size_pt"]*2))})
    return element


def check_contract(path, p):
    require_actions()
    doc = Document(path)
    columns = [int(s._sectPr.find(W+"cols").get(W+"num", "1")) for s in doc.sections]
    expected = [1, p.get("body_columns",2)] if p["english_abstract_position"] == "front" else [1, p.get("body_columns",2), 1]
    if columns != expected: raise StudioError("TEMPLATE_SECTIONS_CHANGED: "+str(columns))
    active = False; bookmark = None
    for item in doc._element.body:
        if item.tag == W+"bookmarkStart" and item.get(W+"name") == "RPS_BODY":
            active = True; bookmark = item.get(W+"id"); continue
        if item.tag == W+"bookmarkEnd" and item.get(W+"id") == bookmark: active = False
        if active and list(item.iter(W+"sectPr")): raise StudioError("BODY_SECTION_FORBIDDEN")
        if active:
            for par in item.iter(W+"p"):
                props=par.find(W+"pPr")
                if props is None or props.find(W+"pStyle") is None or props.find(W+"pStyle").get(W+"val") not in {"RPS"+r for r in ROLES}:
                    raise StudioError("CONTENT_STYLE_FORBIDDEN")
                if any(c.tag not in {W+"pStyle",W+"numPr",W+"jc",W+"keepNext"} for c in props): raise StudioError("CONTENT_PARAGRAPH_LAYOUT_FORBIDDEN")
            for run in item.iter(W+"r"):
                props=run.find(W+"rPr")
                if props is not None and any(c.tag not in {W+"b",W+"i",W+"vertAlign"} for c in props): raise StudioError("CONTENT_RUN_LAYOUT_FORBIDDEN")
    for section in doc.sections:
        margins = section._sectPr.find(W+"pgMar")
        if any(abs(int(margins.get(W+side))-round(p["page"]["margin_cm"][side]*1440/2.54))>1 for side in ("top", "bottom", "left", "right")):
            raise StudioError("TEMPLATE_MARGINS_CHANGED")
    return {"section_columns": columns, "body_columns": p.get("body_columns",2), "body_section_changes": 0, "content_layout_overrides": False}


def import_resources(source, target, elements):
    """Only images and hyperlinks enter the mother; no styles/themes/headers."""
    mapping = {}
    for element in elements:
        for node in element.iter():
            for key, identifier in list(node.attrib.items()):
                if not key.startswith(R): continue
                if identifier not in mapping:
                    relation = source.part.rels[identifier]
                    if relation.reltype == RT.IMAGE and not relation.is_external:
                        # Pandoc already packaged the image. Copy its original bytes
                        # and MIME type without re-decoding it with the narrower
                        # python-docx image-header detector. Drawing extent stays
                        # in the checked content XML and is capped below.
                        from docx.parts.image import ImagePart
                        from docx.opc.packuri import PackURI
                        incoming=relation.target_part
                        collection=target.part.package.image_parts
                        matched=next((part for part in collection if part.blob==incoming.blob and part.content_type==incoming.content_type),None)
                        if matched is None:
                            names={str(part.partname) for part in target.part.package.parts}
                            index=1
                            while "/word/media/rps-image"+str(index)+"."+incoming.partname.ext in names:index+=1
                            partname=PackURI("/word/media/rps-image"+str(index)+"."+incoming.partname.ext)
                            matched=ImagePart(partname,incoming.content_type,incoming.blob)
                            collection.append(matched)
                        mapping[identifier]=target.part.relate_to(matched,RT.IMAGE)
                    elif relation.reltype == RT.HYPERLINK and relation.is_external:
                        mapping[identifier] = target.part.relate_to(relation.target_ref, RT.HYPERLINK, is_external=True)
                    else: raise StudioError("UNSUPPORTED_CONTENT_RELATION: "+relation.reltype)
                node.set(key, mapping[identifier])


def import_lists(source, target, elements):
    """Copy list semantics and restart values, replacing every indent/font."""
    numbering = target.part.numbering_part.element
    originals = source.part.numbering_part.element
    mapping = {}
    for element in elements:
        for reference in element.iter(W+"numId"):
            old = reference.get(W+"val")
            if old not in mapping:
                num = next((n for n in originals.findall(W+"num") if n.get(W+"numId") == old), None)
                if num is None: raise StudioError("INVALID_LIST_NUMBERING")
                abstract_id = num.find(W+"abstractNumId").get(W+"val")
                abstract = next(n for n in originals.findall(W+"abstractNum") if n.get(W+"abstractNumId") == abstract_id)
                aid = max([int(n.get(W+"abstractNumId")) for n in numbering.findall(W+"abstractNum")]+[0])+1
                nid = max([int(n.get(W+"numId")) for n in numbering.findall(W+"num")]+[0])+1
                clean = OxmlElement("w:abstractNum"); clean.set(W+"abstractNumId", str(aid))
                for level in abstract.findall(W+"lvl"):
                    item = OxmlElement("w:lvl"); item.set(W+"ilvl", level.get(W+"ilvl"))
                    for tag in ("start", "numFmt", "lvlText"):
                        value = level.find(W+tag)
                        if value is not None: item.append(deepcopy(value))
                    props = OxmlElement("w:pPr")
                    set_child(props, "ind", {"w:left": str(400*(int(level.get(W+"ilvl"))+1)), "w:hanging": "400"})
                    item.append(props); clean.append(item)
                numbering.append(clean)
                copy = deepcopy(num); copy.set(W+"numId", str(nid)); copy.find(W+"abstractNumId").set(W+"val", str(aid))
                for child in list(copy):
                    if child.tag not in {W+"abstractNumId", W+"lvlOverride"}: copy.remove(child)
                numbering.append(copy); mapping[old] = str(nid)
            reference.set(W+"val", mapping[old])


def fill(native, metadata, equations, master, p, output):
    require_actions()
    from .math_layout import wrap, tree
    from .docx_layout import move_table_captions
    source = Document(native)
    from docxtpl import DocxTemplate
    from jinja2 import Environment, StrictUndefined
    import tempfile
    context={key+"_"+lang:"" for lang in ("zh","en") for key in ("title","authors","affiliation","abstract","keywords")}
    context["RPS_BODY"]="{{RPS_BODY}}"
    for key,value in metadata.items():
        if key not in context: continue
        if isinstance(value,list): value="; ".join(value)
        if any(c in str(value) for c in ("\f","\a")): raise StudioError("CONTENT_CONTROL_CHARACTER")
        if value and key.startswith("abstract_"): value=("摘要：" if key.endswith("zh") else "Abstract: ")+str(value)
        if value and key.startswith("keywords_"): value=("关键词：" if key.endswith("zh") else "Keywords: ")+str(value)
        context[key]=str(value)
    tpl=DocxTemplate(master); tpl.render(context,jinja_env=Environment(undefined=StrictUndefined),autoescape=True)
    rendered=BytesIO(); tpl.save(rendered); rendered.seek(0)
    doc=Document(rendered); body = doc._element.body
    if metadata.get("_meta_docx"):
        meta_source=Document(metadata["_meta_docx"])
        meta_elements=[deepcopy(e) for e in meta_source._element.body if e.tag!=W+"sectPr"]
        import_resources(meta_source,doc,meta_elements)
        if len(meta_elements)!=len(metadata["_meta_roles"]): raise StudioError("METADATA_OBJECT_COUNT_CHANGED")
        for key,element in zip(metadata["_meta_roles"],meta_elements):
            role="Title" if key.startswith("title") else "Author" if key.startswith(("authors","affiliation")) else "Abstract" if key.startswith("abstract") else "Keywords"
            target=next((e for e in body if e.tag==W+"p" and "".join(t.text or "" for t in e.iter(W+"t"))==context[key] and e.find(W+"pPr/"+W+"pStyle") is not None and e.find(W+"pPr/"+W+"pStyle").get(W+"val")=="RPS"+role),None)
            if target is None:
                if p.get("anonymous") and role=="Author" or key[-2:] not in p.get("metadata_languages",["zh","en"]): continue
                raise StudioError("METADATA_ROLE_MISSING: "+key)
            clean_content(element,p); pstyle(element,role); target.addprevious(element); body.remove(target)
    for element in list(body):
        if element.tag==W+"p" and not "".join(t.text or "" for t in element.iter(W+"t")) and not list(element.iter(M+"oMath")) and element.find(W+"pPr/"+W+"sectPr") is None:
            body.remove(element)
    original = [clean_content(deepcopy(e), p) for e in source._element.body if e.tag != W+"sectPr"]
    import_resources(source, doc, original); import_lists(source, doc, original)
    for element in list(body):
        text = "".join(t.text or "" for t in element.iter(W+"t"))
        match = re.fullmatch(r"\{\{(\w+)\}\}", text)
        if not match or match[1] == "RPS_BODY": continue
        key = match[1]; value = metadata.get(key, "")
        if isinstance(value, list): value = "; ".join(value)
        if not value: body.remove(element); continue
        if key.startswith("abstract_"): value = ("摘要：" if key.endswith("zh") else "Abstract: ")+str(value)
        if key.startswith("keywords_"): value = ("关键词：" if key.endswith("zh") else "Keywords: ")+str(value)
        element.find(W+"r/"+W+"t").text = str(value)
    anchor = next(e for e in body if "".join(t.text or "" for t in e.iter(W+"t")) == "{{RPS_BODY}}")
    start = OxmlElement("w:bookmarkStart"); start.set(W+"id", "9181"); start.set(W+"name", "RPS_BODY"); anchor.addprevious(start)
    width = column_width(p); ledger = []; inline_ledger = []; index = 0
    for element in original:
        if element.tag == W+"p" and list(element.iter(M+"oMathPara")):
            if index >= len(equations): raise StudioError("FORMULA_COUNT_CHANGED")
            equation = equations[index]; index += 1
            formulae = list(element.iter(M+"oMath"))
            if len(formulae) != 1: raise StudioError("UNSUPPORTED_DISPLAY_STRUCTURE: "+equation["id"])
            # Keep an extra 10 pt safety allowance for renderer width
            # differences; isolation also checks the actual number boundary.
            changes = wrap(formulae[0], (width-550)/20-26,p["body_size_pt"])
            ledger.append({"id": equation["id"], "changes": changes})
            table = doc.add_table(rows=1, cols=2); table.autofit = False
            table._tbl.getparent().remove(table._tbl)
            set_child(table._tbl.tblPr, "tblW", {"w:w": str(width), "w:type": "dxa"})
            set_child(table._tbl.tblPr, "tblDescription", {"w:val": "RPS_EQUATION_LAYOUT"})
            borders = set_child(table._tbl.tblPr, "tblBorders")
            for edge in ("top", "bottom", "left", "right", "insideH", "insideV"): set_child(borders, edge, {"w:val": "nil"})
            margins = set_child(table._tbl.tblPr, "tblCellMar")
            for edge in ("top", "bottom", "left", "right"): set_child(margins, edge, {"w:w": "0", "w:type": "dxa"})
            for cell, column, cell_width in zip(table.rows[0].cells, table.columns, (width-550, 550)):
                cell.width = Twips(cell_width); column.width = Twips(cell_width)
                set_child(cell._tc.get_or_add_tcPr(),"vAlign",{"w:val":"center"})
            formula = table.cell(0, 0).paragraphs[0]._p; pstyle(formula, "Equation")
            # The native display wrapper stays editable, including original cases.
            for child in list(element):
                if child.tag == M+"oMathPara": formula.append(child)
            label = equation["id"] if p["equation_numbering"] == "preserve" else str(index)
            right = table.cell(0, 1).paragraphs[0]._p; pstyle(right, "Equation")
            right.append(paragraph("("+label+")", "Equation").find(W+"r"))
            set_child(right.find(W+"pPr"), "jc", {"w:val": "right"})
            set_child(table._tbl.find(W+"tr").get_or_add_trPr(), "cantSplit")
            anchor.addprevious(table._tbl)
        else:
            # Inline formulas in table cells use that cell's column budget.
            for formula in element.iter(M+"oMath"):
                cell = next((a for a in formula.iterancestors() if a.tag == W+"tc"), None)
                cell_budget = width/20-12
                if cell is not None:
                    table = next(a for a in cell.iterancestors() if a.tag == W+"tbl")
                    cells = len(table.find(W+"tr").findall(W+"tc"))
                    cell_budget = width/20/cells-12
                changes = wrap(formula, cell_budget,p["body_size_pt"])
                inline_ledger.append(changes)
            anchor.addprevious(element)
    if index != len(equations): raise StudioError("FORMULA_COUNT_CHANGED")
    end = OxmlElement("w:bookmarkEnd"); end.set(W+"id", "9181"); anchor.addprevious(end); body.remove(anchor)
    moved = 0
    for table in doc.tables:
        if is_equation_table(table._tbl): continue
        if len(table.columns) == 0: raise StudioError("EMPTY_TABLE")
        props = table._tbl.tblPr
        for child in list(props): props.remove(child)
        table.autofit = False
        set_child(props, "tblW", {"w:w": str(width), "w:type": "dxa"})
        for column in table.columns: column.width = Twips(width//len(table.columns))
        for row in table.rows:
            for cell in row.cells:
                cell.width = Twips(width//len(table.columns))
                for child in list(cell._tc.get_or_add_tcPr()):
                    if child.tag not in {W+"tcW", W+"gridSpan", W+"vMerge"}: cell._tc.tcPr.remove(child)
                for par in cell.paragraphs: pstyle(par._p, "TableText")
        _three_line(table, p["tables"], None)
    for shape in doc.inline_shapes:
        maximum = width*635
        if shape.width > maximum:
            ratio = maximum/shape.width; shape.height = round(shape.height*ratio); shape.width = maximum
    for element in body:
        if element.tag==W+"p" and list(element.iter(W+"drawing")):
            pstyle(element,"Equation")
    if p.get("captions"):
        from .captions import arrange
        try: moved=arrange(doc,p)
        except StudioError:
            # Keep the actual first diagnostic Word without claiming acceptance.
            doc.save(output)
            write_json(Path(output).parent/"math-layout-ledger.json",ledger)
            write_json(Path(output).parent/"inline-math-layout-ledger.json",inline_ledger)
            raise
    else: moved=move_table_captions(doc)
    if list(body.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}anchor")):
        raise StudioError("FLOATING_IMAGE_FORBIDDEN")
    doc.save(output)
    with zipfile.ZipFile(master) as original_package, zipfile.ZipFile(output) as final_package:
        for part in ("word/styles.xml","word/settings.xml"):
            if original_package.read(part) != final_package.read(part): raise StudioError("TEMPLATE_PART_CHANGED: "+part)
    write_json(Path(output).parent/"math-layout-ledger.json", ledger)
    write_json(Path(output).parent/"inline-math-layout-ledger.json", inline_ledger)
    return {**check_contract(output, p), "display_count": index, "table_captions_moved": moved, "math_layout_changes": sum(len(e["changes"]) for e in ledger), "native_math_count": len(list(body.iter(M+"oMath")))}

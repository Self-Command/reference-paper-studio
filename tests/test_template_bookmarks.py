"""Template engines may renumber bookmarks while preserving their names."""
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from studio.docx_layout import body_only_docx


def test_body_audit_follows_remapped_bookmark_and_excludes_english_end(tmp_path):
    doc = Document()
    doc.add_paragraph("首页摘要")
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "0")
    start.set(qn("w:name"), "RPS_BODY")
    doc._element.body.insert(len(doc._element.body)-1, start)
    doc.add_paragraph("正文数值 12.50")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "0")
    doc._element.body.insert(len(doc._element.body)-1, end)
    doc.add_paragraph("End English abstract")
    source, output = tmp_path / "filled.docx", tmp_path / "audit.docx"
    doc.save(source)
    body_only_docx(source, output)
    assert [p.text for p in Document(output).paragraphs] == ["正文数值 12.50"]

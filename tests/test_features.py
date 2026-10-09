import json
from copy import deepcopy
from pathlib import Path
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from studio.runtime import StudioError, digest
from studio.references import resolve_rules, inspect_pdf
from studio.manuscript import parse, validate_assets, validate_references, model_gate, resolve_crossrefs
from studio.docx_layout import move_table_captions, apply_layout
from studio.pipeline import strict_pandoc
from studio.qa import check_fonts, check_docx, assert_text, check_equation_visibility, latex_reader_copy, check_page_flow
from studio.vendor.latex_to_omml import latex_to_omml

ROOT = Path(__file__).resolve().parents[1]
PROFILE = json.loads((ROOT/"assets/profiles/chinese-journal.json").read_text())


def observation(value, source):
    return {"property":"layout", "value":value, "source":source, "pages":[1], "reviewed":True}


def test_reference_consensus_beats_official():
    result = resolve_rules([observation(2,"P1"), observation(2,"P2")], {"layout":{**observation(1,"J1")}})
    assert result[0]["value"] == 2
    assert result[0]["status"] == "reference_consensus"


def test_conflict_uses_official_only_with_evidence():
    assert resolve_rules([observation(2,"P1"),observation(1,"P2")])[0]["status"] == "conflicting"
    assert resolve_rules([observation(2,"P1"),observation(1,"P2")], {"layout":observation(1,"J1")})[0]["status"] == "official_disambiguation"


def test_single_source_and_bad_extraction_unconfirmed():
    assert resolve_rules([observation(2,"P1")])[0]["status"] == "unconfirmed"
    assert resolve_rules([{**observation(2,"P1"),"quality":"needs_review"},observation(2,"P2")])[0]["status"] == "unconfirmed"


def test_scanned_pdf_unreliable(tmp_path):
    from pypdf import PdfWriter
    p = tmp_path/"scan.pdf"; w = PdfWriter(); w.add_blank_page(width=595,height=842); w.write(p)
    assert inspect_pdf(p)["pages"][0]["quality"] == "needs_review"


def test_original_model_has_all_30_equations(tmp_path):
    _, eq = parse(ROOT/"fixtures/model/model.md", tmp_path)
    assert [e["id"] for e in eq] == [f"M{i}" for i in range(1,31)]


def test_duplicate_equation_rejected(tmp_path):
    p=tmp_path/"dup.md"; p.write_text(r"\[x\tag{M1}\]"+"\n\n"+r"\[y\tag{M1}\]")
    with pytest.raises(StudioError, match="DUPLICATE_EQUATION"): parse(p,tmp_path)


def test_missing_asset_rejected():
    with pytest.raises(StudioError, match="MISSING_IMAGE"):
        validate_assets({"t":"Image","c":[[],[],["missing.png",""]]},ROOT)


def test_broken_numeric_reference():
    with pytest.raises(StudioError, match="BROKEN_REFERENCE"):
        validate_references({"blocks":[{"t":"Str","c":"正文[3]参考文献[1]"}]})


def test_unverified_model_blocked():
    with pytest.raises(StudioError, match="MODEL_NOT_VERIFIED"): model_gate(ROOT,"fixtures/model/model.md",None,True)


def test_stale_verification_blocked(tmp_path):
    (tmp_path/"model.md").write_text("model")
    (tmp_path/"report.md").write_text("human verification report")
    (tmp_path/"evidence.json").write_text(json.dumps({"status":"passed","author_approved":True,"model":"model.md","report":"report.md","model_sha256":"old","report_sha256":digest(tmp_path/"report.md")}))
    with pytest.raises(StudioError, match="MODEL_NOT_VERIFIED"): model_gate(tmp_path,"model.md","evidence.json",True)


def test_missing_experiment_source_blocked(tmp_path):
    for name in ("model.md","report.md"): (tmp_path/name).write_text(name)
    evidence={"status":"passed","author_approved":True,"model":"model.md","report":"report.md","model_sha256":digest(tmp_path/"model.md"),"report_sha256":digest(tmp_path/"report.md"),"quantitative_claims":[{"value":"97%"}]}
    (tmp_path/"evidence.json").write_text(json.dumps(evidence))
    with pytest.raises(StudioError, match="CLAIM_WITHOUT_EVIDENCE"): model_gate(tmp_path,"model.md","evidence.json",True)


def test_unknown_formula_command_fails(tmp_path):
    p=tmp_path/"bad.md"; p.write_text(r"$\unsupportedstudio{x}$")
    with pytest.raises(StudioError, match="CONVERSION_FAILED"): strict_pandoc(["pandoc",p,"-o",tmp_path/"bad.docx"])


def test_mit_formula_converter_native():
    result=latex_to_omml(r"\sum_{i=1}^n \frac{x_i^2}{y_i}")
    assert result is not None


def test_caption_repair_does_not_treat_equation_as_table(tmp_path):
    doc=Document(); table=doc.add_table(rows=2,cols=2)
    table.cell(0,0).text="项目"; table.cell(0,1).text="数值"
    doc.add_paragraph("表1 示例表 / Table 1 Example")
    doc.add_paragraph("公式另行处理")
    assert move_table_captions(doc)==1
    assert table._tbl.getprevious().tag == qn("w:p")


@pytest.mark.parametrize("position", ["front","end"])
def test_front_body_sections_preserved(tmp_path,position):
    doc=Document(); doc.add_paragraph("模型正文"); path=tmp_path/"layout.docx"; doc.save(path)
    profile=deepcopy(PROFILE); profile["english_abstract_position"]=position
    apply_layout(path,{"title_zh":"中文题名","abstract_zh":"摘要","abstract_en":"English abstract"},profile,[])
    assert check_docx(path)["section_columns"][0]==1
    assert 2 in check_docx(path)["section_columns"]
    final=Document(path)
    assert final.sections[-1]._sectPr.find(qn("w:type")).get(qn("w:val"))=="continuous"


def test_missing_font_fails():
    with pytest.raises(StudioError, match="MISSING_FONT"): check_fonts({"fonts":{"cjk":"No Such Studio Font"}})


def test_text_drift_fails():
    with pytest.raises(StudioError, match="CONTENT_CHANGED"):
        assert_text({"blocks":[{"t":"Str","c":"123"}]},{"blocks":[{"t":"Str","c":"124"}]},"test")


def test_crossrefs_share_numbering():
    ast={"blocks":[{"t":"Str","c":"{{eq:M1}}"}]}
    assert resolve_crossrefs(ast,[{"id":"M1","ordinal":1}],"sequential")["eq:M1"]=="(1)"
    assert ast["blocks"][0]["c"]=="(1)"


def test_broken_crossref_fails():
    with pytest.raises(StudioError,match="BROKEN_CROSS_REFERENCE"):
        resolve_crossrefs({"blocks":[{"t":"Str","c":"{{eq:Missing}}"}]},[],"preserve")


@pytest.mark.parametrize("language,expected", [("zh", "图1 表1 (M1)"), ("en", "Figure 1 Table 1 (M1)")])
def test_generated_object_references_follow_registered_language(language, expected):
    ast={"blocks":[
        {"t":"Figure","c":[["fig:sample",[],[]],[None,[]],[]]},
        {"t":"Table","c":[["tab:data",[],[]],[None,[]],[],[],[],[]]},
        {"t":"Str","c":"{{fig:sample}} {{tab:data}} {{eq:M1}}"},
        {"t":"Str","c":"Authored 图9 and Table 8 are untouched."}
    ]}
    mapping=resolve_crossrefs(ast,[{"id":"M1","ordinal":1}],"preserve",language)
    assert ast["blocks"][2]["c"] == expected
    assert ast["blocks"][3]["c"] == "Authored 图9 and Table 8 are untouched."
    assert mapping["eq:M1"] == "(M1)"


def test_quoted_prose_is_preserved_in_semantic_check():
    from studio.manuscript import text_of
    assert text_of({"t":"Quoted","c":[{"t":"DoubleQuote"},[{"t":"Str","c":"时移0"}]]})=="“时移0”"


def test_results_section_needs_source_data(tmp_path):
    (tmp_path/"paper.md").write_text("# 实验结果\n效率提高97%",encoding="utf-8")
    (tmp_path/"report.md").write_text("review")
    evidence={"status":"passed","author_approved":True,"model":"paper.md","report":"report.md","model_sha256":digest(tmp_path/"paper.md"),"report_sha256":digest(tmp_path/"report.md")}
    (tmp_path/"evidence.json").write_text(json.dumps(evidence))
    with pytest.raises(StudioError,match="RESULTS_WITHOUT_DATA"): model_gate(tmp_path,"paper.md","evidence.json",True)


def test_real_invisible_math_pdf_regression():
    with pytest.raises(StudioError,match="WORD_PDF_MATH_NOT_VISIBLE"):
        check_equation_visibility(ROOT/"fixtures/regressions/invisible-math.pdf",[f"M{i}" for i in range(1,31)])


def test_only_identical_repeated_page_headers_normalized(tmp_path):
    source=tmp_path/"tables.tex"
    first=r"\begin{longtable}{ll}\toprule A&B\\\endhead 1&2\\\end{longtable}"
    second=r"\begin{longtable}{ll}\toprule X&Y\\\midrule\endfirsthead\toprule X&Y\\\midrule\endhead 3&4\\\end{longtable}"
    source.write_text(first+second)
    result=latex_reader_copy(source,tmp_path).read_text()
    assert result.count("X&Y")==1 and "A&B" in result and "3&4" in result
    source.write_text(second.replace(r"\endfirsthead\toprule X&Y",r"\endfirsthead\toprule X&Z"))
    with pytest.raises(StudioError,match="UNEXPECTED_REPEATED_TABLE_HEADER"): latex_reader_copy(source,tmp_path)


def test_citeproc_numeric_and_missing_keys(tmp_path):
    source=tmp_path/"cite.md"; source.write_text("A reference [@yuan2007].")
    bibliography=tmp_path/"references.json"
    bibliography.write_text(json.dumps([{"id":"yuan2007","type":"paper-conference","title":"Allocating dynamic time-spectrum blocks in cognitive radio networks","author":[{"family":"Yuan","given":"Y"}],"issued":{"date-parts":[[2007]]}}]))
    result=strict_pandoc(["pandoc",source,"-t","plain","--citeproc","--bibliography",bibliography,"--csl",ROOT/"assets/numbered.csl"])
    assert "[1]" in result and "Allocating dynamic" in result
    source.write_text("Missing [@unknown-studio-key].")
    with pytest.raises(StudioError,match="CONVERSION_FAILED"):
        strict_pandoc(["pandoc",source,"-t","plain","--citeproc","--bibliography",bibliography])


def test_equation_layout_never_receives_data_table_borders(tmp_path):
    from studio.docx_layout import configure_math_font, is_equation_table
    source=tmp_path/"native.md";source.write_text(r"$$\frac{x_i}{1+x_i}$$")
    path=tmp_path/"native.docx"
    strict_pandoc(["pandoc",source,"-o",path])
    configure_math_font(path,PROFILE)
    apply_layout(path,{},PROFILE,[{"id":"M1","source":r"\frac{x_i}{1+x_i}","ordinal":1}])
    doc=Document(path)
    assert len(doc.tables)==0
    assert check_docx(path)["data_tables"]==0
    assert check_docx(path)["equation_layouts"]==0
    assert "(M1)" in "".join(p.text for p in doc.paragraphs)


def test_short_equations_do_not_fragment_body_sections(tmp_path):
    source=tmp_path/"short.md"; source.write_text("模型正文。\n\n"+r"$$x_i+y_i=z_i$$"+"\n\n后续正文。")
    path=tmp_path/"short.docx"; strict_pandoc(["pandoc",source,"-o",path])
    report=apply_layout(path,{},PROFILE,[{"id":"E1","ordinal":1}])
    assert report["equation_columns"][0]["columns_spanned"]==1
    assert len(Document(path).sections)==2


def test_blank_interior_page_fails_flow_check(tmp_path):
    from pypdf import PdfWriter
    path=tmp_path/"pages.pdf"; writer=PdfWriter()
    writer.add_blank_page(width=595,height=842); writer.write(path)
    # A short final page is allowed; a blank page before remaining content is not.
    assert check_page_flow(path,PROFILE)["pages"]==1
    writer.add_blank_page(width=595,height=842); writer.write(path)
    with pytest.raises(StudioError,match="SPARSE_INTERIOR_PAGE"): check_page_flow(path,PROFILE)


@pytest.mark.parametrize("body_overflow",[False,True])
def test_abstract_repeated_in_body_does_not_change_column_region(monkeypatch,body_overflow):
    from types import SimpleNamespace
    from studio import qa
    def line(text,top,x=63):
        return [{"text":letter,"x0":x+i*5,"x1":x+(i+1)*5,"top":top,"bottom":top+10} for i,letter in enumerate(text)]
    chars=line("Repeated text entry",100)+line("metadata-end-anchor",200)+line("Repeated text entry",300)
    chars += [{"text":"X","x0":296,"x1":302,"top":120,"bottom":130}]
    if body_overflow: chars += [{"text":"Y","x0":296,"x1":302,"top":350,"bottom":360}]
    page=SimpleNamespace(width=595.3,height=841.9,chars=chars)
    class Proof:
        pages=[page]
        def __enter__(self): return self
        def __exit__(self,*args): pass
    monkeypatch.setattr(qa.pdfplumber,"open",lambda path:Proof())
    ast={"blocks":[{"t":"Para","c":[{"t":"Str","c":"Repeated text entry"}]}]}
    result=qa.check_column_geometry("cloud-proof.pdf",PROFILE,ast,{"abstract_zh":"Repeated text entry","keywords_en":["metadata-end-anchor"]})
    assert result["regions"][0]["body_top"]==299
    assert [v["text"] for v in result["violations"]]==(["Y"] if body_overflow else [])

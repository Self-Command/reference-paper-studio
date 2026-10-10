from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
import pytest
from studio.composition import load_composition
from studio.components import compile_components
from studio.locked_template import create_templates,check_contract
from studio.runtime import StudioError,run

def test_page_number_component_controls_real_footer_and_rendered_sequence(tmp_path):
    _,p=load_composition("reference-journal-v1","2.0.0")
    create_templates(tmp_path,p)
    path=tmp_path/"master-front.docx"
    doc=Document(path)
    # Add independent pages solely to test rendered field updates. Mother
    # pages/sections are otherwise untouched; all of this executes in Actions.
    doc.add_paragraph("First page")
    doc.add_page_break();doc.add_paragraph("Second page")
    doc.save(path)
    assert all([field.get(qn("w:instr")) for field in section.footer._element.iter(qn("w:fldSimple"))]==["PAGE"] for section in doc.sections)
    run(["libreoffice","-env:UserInstallation=file:///tmp/rps-footer-test","--headless","--convert-to","pdf","--outdir",tmp_path,path])
    import pdfplumber
    with pdfplumber.open(tmp_path/"master-front.pdf") as pdf:
        assert len(pdf.pages)==2
        for index,page in enumerate(pdf.pages,1):
            footer=[word["text"] for word in page.extract_words() if word["top"]>page.height-60]
            assert footer==[str(index)]
    p["english_abstract_position"]="front"
    doc.sections[0].footer.paragraphs[0].text="changed"
    doc.save(path)
    with pytest.raises(StudioError,match="TEMPLATE_PAGE_NUMBER_CHANGED"):check_contract(path,p)

def test_component_removal_does_not_leave_hidden_page_fields(tmp_path):
    descriptor=deepcopy(load_composition("reference-journal-v1","2.0.0")[0])
    descriptor["components"].remove("page-number")
    p,_=compile_components(descriptor)
    p.update(id=descriptor["id"],version=descriptor["version"])
    create_templates(tmp_path,p)
    assert not any(list(section.footer._element.iter(qn("w:fldSimple"))) for section in Document(tmp_path/"master-front.docx").sections)


@pytest.mark.parametrize("variant",["front","end"])
def test_single_column_front_starts_body_on_next_page_and_renders_title(tmp_path,variant):
    from studio.locked_template import fill
    from studio.frontmatter import prepare
    from studio.qa import check_rendered_frontmatter,check_page_flow
    _,p=load_composition("gmcm","1.1.0")
    p["english_abstract_position"]=variant
    create_templates(tmp_path/"master",p)
    native=tmp_path/"native.docx";doc=Document();doc.add_heading("BodyStart",1);doc.add_paragraph("Actual body content.");doc.save(native)
    metadata=prepare({"title_zh":"中文题目可见","abstract_zh":"独立摘要页","keywords_zh":["测试"]},[1,23,1],tmp_path)
    result=tmp_path/"manuscript.docx"
    fill(native,metadata,[],tmp_path/"master"/("master-"+variant+".docx"),p,result)
    rendered=Document(result)
    assert rendered.sections[0]._sectPr.find(qn("w:type")).get(qn("w:val"))=="continuous"
    assert rendered.sections[1]._sectPr.find(qn("w:type")).get(qn("w:val"))=="nextPage"
    run(["libreoffice","-env:UserInstallation=file:///tmp/rps-front-test-"+variant,"--headless","--convert-to","pdf","--outdir",tmp_path,result])
    ast={"blocks":[{"t":"Header","c":[1,["",[],[]],[{"t":"Str","c":"BodyStart"}]]}]}
    proof=tmp_path/"manuscript.pdf"
    observed=check_rendered_frontmatter(proof,p,ast,metadata,tmp_path/"metadata-source.json")
    assert observed["title_visibility"]==[{"role":"title_zh","pages":[1],"math_visibility":"checked separately"}]
    assert check_page_flow(proof,p)["pages"]==2


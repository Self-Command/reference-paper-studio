from pathlib import Path
import json
from copy import deepcopy
import pytest
from docx import Document
from docx.oxml.ns import qn
from studio.runtime import StudioError,digest
from studio.registry import build,resolve
from studio.composition import load_composition
from studio.locked_template import check_contract,column_width,fill
from studio.captions import arrange

def test_single_column_template_budget_and_contract(tmp_path,monkeypatch):
    monkeypatch.setenv("RPS_REGISTRY_ROOT",str(tmp_path))
    target=build("cumcm","1.0.0")
    p=load_composition("cumcm","1.0.0")[1]
    p["english_abstract_position"]="front"
    assert column_width(p)>9000
    assert check_contract(target/"master-front.docx",p)["section_columns"]==[1,1]

def test_candidates_never_resolve_as_releases(tmp_path,monkeypatch):
    monkeypatch.setenv("RPS_REGISTRY_ROOT",str(tmp_path))
    target=build("reference-journal-v1","2.0.0")
    request={"template":{"id":"reference-journal-v1","version":"2.0.0","variant":"front","font_profile":"open-serif","manifest_sha256":digest(target/"manifest.json")}}
    with pytest.raises(StudioError,match="TEMPLATE_NOT_READY"): resolve(request)
    monkeypatch.setenv("RPS_CANDIDATE_BUILD","true")
    assert resolve(request)[0].is_file()
    with (target/"master-front.docx").open("ab") as stream: stream.write(b"tampered")
    with pytest.raises(StudioError,match="TEMPLATE_INTEGRITY"): resolve(request)

def test_data_caption_moves_below_and_does_not_gain_vertical_borders():
    doc=Document(); caption=doc.add_paragraph("表1 参数 / Table 1 Parameters"); table=doc.add_table(rows=2,cols=2)
    for row in table.rows:
        for cell in row.cells: cell.text="value"
    from studio.locked_template import pstyle
    for par in doc.paragraphs: pstyle(par._p,"Caption")
    for cell in table.rows[-1].cells:
        for par in cell.paragraphs: pstyle(par._p,"TableText")
    assert arrange(doc,{"captions":{"figure":"above","table":"below"}})==1
    assert table._tbl.getnext() is caption._p
    assert table.rows[-1].cells[0].paragraphs[0]._p.find(qn("w:pPr")+"/"+qn("w:keepNext")) is not None

def test_contest_year_and_ieee_identity_cannot_be_substituted(tmp_path,monkeypatch):
    monkeypatch.setenv("RPS_REGISTRY_ROOT",str(tmp_path))
    monkeypatch.setenv("RPS_CANDIDATE_BUILD","true")
    target=build("gmcm","1.1.0")
    identity={"id":"gmcm","version":"1.1.0","year":2026,"variant":"front","font_profile":"open-serif","manifest_sha256":digest(target/"manifest.json")}
    assert resolve({"template":identity})[0].is_file()
    with pytest.raises(StudioError,match="TEMPLATE_YEAR_MISMATCH"):
        resolve({"template":{**identity,"year":2025}})
    with pytest.raises(StudioError,match="UNKNOWN_TEMPLATE"):
        load_composition("ieee-generic","1.0.0")


def test_template_build_reuses_frozen_asset_and_never_resets_release(tmp_path,monkeypatch):
    monkeypatch.setenv("RPS_REGISTRY_ROOT",str(tmp_path))
    target=build("reference-journal-v1","2.0.0")
    manifest=json.loads((target/"manifest.json").read_text())
    # Synthetic maintenance guard fixture, not a suite release declaration.
    manifest.update(release_state="released",automatic_pass=True,visual_review_complete=True)
    (target/"manifest.json").write_text(json.dumps(manifest))
    before={path.name:digest(path) for path in target.iterdir() if path.is_file()}
    assert build("reference-journal-v1","2.0.0")==target
    assert before=={path.name:digest(path) for path in target.iterdir() if path.is_file()}
    changed=deepcopy(manifest);changed["composition_sha256"]="0"*64
    (target/"manifest.json").write_text(json.dumps(changed))
    fingerprint=digest(target/"manifest.json")
    with pytest.raises(StudioError,match="TEMPLATE_REGISTRY_IDENTITY_CONFLICT"):
        build("reference-journal-v1","2.0.0")
    assert digest(target/"manifest.json")==fingerprint


def test_build_cannot_repair_tampered_mother_silently(tmp_path,monkeypatch):
    monkeypatch.setenv("RPS_REGISTRY_ROOT",str(tmp_path))
    target=build("reference-journal-v1","2.0.0")
    with (target/"master-front.docx").open("ab") as stream:stream.write(b"tampered")
    fingerprint=digest(target/"master-front.docx")
    with pytest.raises(StudioError,match="TEMPLATE_INTEGRITY_MISMATCH"):
        build("reference-journal-v1","2.0.0")
    assert digest(target/"master-front.docx")==fingerprint


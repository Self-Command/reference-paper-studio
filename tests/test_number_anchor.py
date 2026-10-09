from pathlib import Path
import json
from docx import Document
from studio.manuscript import parse
from studio.pipeline import strict_pandoc
from studio.locked_template import create_templates,fill,policy
from studio.docx_layout import body_only_docx,M,W
from studio.math_layout import tree

def test_single_line_has_shared_baseline_and_audit_restores_display_wrapper(tmp_path):
    source=tmp_path/"source.md"
    source.write_text(r"$$\frac{x_i^2}{1+x_i}=y_i.\tag{E1}$$"+"\n\n"+r"$$\begin{aligned}a&=b\\c&=d\end{aligned}\tag{E2}$$")
    ast,equations=parse(source,tmp_path)
    native=tmp_path/"native.docx";strict_pandoc(["pandoc",tmp_path/"document.json","-f","json","-o",native])
    create_templates(tmp_path/"master")
    output=tmp_path/"manuscript.docx";fill(native,{},equations,tmp_path/"master/master-front.docx",policy(),output)
    doc=Document(output); first,second=doc.tables
    assert len(first.columns)==1 and len(second.columns)==2
    paragraph=first.cell(0,0).paragraphs[0]._p
    assert paragraph.find(W+"pPr/"+W+"textAlignment").get(W+"val")=="baseline"
    assert paragraph.find(M+"oMath") is not None and paragraph.find(M+"oMathPara") is None
    assert [item["number_anchor"] for item in json.loads((tmp_path/"math-layout-ledger.json").read_text())]==["baseline","center"]
    audit=tmp_path/"audit.docx";body_only_docx(output,audit)
    result=Document(audit)
    assert len(list(result._element.iter(M+"oMathPara")))==2
    assert [tree(m) for m in result._element.iter(M+"oMath")]==[tree(m) for m in Document(native)._element.iter(M+"oMath")]
    assert "(E1)" not in "".join(p.text for p in result.paragraphs)

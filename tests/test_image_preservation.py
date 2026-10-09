import pytest
from PIL import Image
from docx import Document
from studio.qa import check_image_assets
from studio.runtime import StudioError

def test_image_bytes_and_occurrences_cannot_be_masked_by_alt_text(tmp_path):
    image=tmp_path/"image.png";Image.new("RGB",(20,20),"white").save(image)
    source={"blocks":[{"t":"Image","c":[["",[],[]],[],[image.name,""]]}]}
    doc=Document();doc.add_picture(str(image));word=tmp_path/"word.docx";doc.save(word)
    assert check_image_assets(source,word,source,tmp_path,tmp_path)["count"]==1
    doc=Document();doc.save(word)
    with pytest.raises(StudioError,match="IMAGE_CONTENT_OR_COUNT_CHANGED: Word"):
        check_image_assets(source,word,source,tmp_path,tmp_path)
